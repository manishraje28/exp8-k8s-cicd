import os
from docx import Document
from docx.shared import Inches, Pt

here = os.path.dirname(os.path.abspath(__file__))
S = lambda f: os.path.join(here, 'screenshots', f)
d = Document()
st = d.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11)


def code(txt):
    p = d.add_paragraph(); r = p.add_run(txt); r.font.name = 'Consolas'; r.font.size = Pt(9)


def img(f, cap):
    d.add_picture(S(f), width=Inches(6.3)); c = d.add_paragraph(cap); c.runs[0].italic = True


def bold(t):
    d.add_paragraph().add_run(t).bold = True


d.add_heading('Experiment 8: Deploying Applications, CI/CD Pipeline, and Cluster Monitoring', 0)
d.add_paragraph('Name: Manish Raje\nCourse Outcomes: CO3, CO5, CO6\n'
                'Environment: Windows 11, Docker Desktop 28.3.2, Minikube v1.39.0 (Kubernetes v1.37.0), kubectl v1.32.2, Helm v4.3.1')

d.add_heading('1. GitHub Repository', 1)
d.add_paragraph('https://github.com/manishraje28/exp8-k8s-cicd')
d.add_paragraph('Docker Hub image: https://hub.docker.com/r/manishraje28/sample-web-app')

d.add_heading('2. Pre-Lab Diagnostic Quiz', 1)
qa = [
    ('Q1. What is the structural relationship between a Deployment, a ReplicaSet and a Pod?',
     'A Deployment is the top-level declarative object. It creates and manages ReplicaSets, one per revision of the Pod template. '
     'Each ReplicaSet keeps the requested number of identical Pods running and replaces any that fail. When the Deployment changes '
     '(for example a new image), it creates a new ReplicaSet and scales it up while scaling the old one down, which gives rolling '
     'updates and rollback. In this lab, sample-web-app first owned ReplicaSet 5cf58b94f4 (nginx:alpine); after the image update '
     'a new ReplicaSet 64dbf556ff (the CI-built image) replaced it.'),
    ('Q2. Why should Docker Hub credentials never be stored in .github/workflows/ci-cd.yml? Where should they be stored?',
     'The workflow file is committed to the repository, so anything in it is visible to anyone who can read the repo and stays in '
     'Git history even after it is removed. A leaked token lets anyone push or overwrite images. Credentials belong in GitHub '
     'encrypted Repository Secrets (Settings > Secrets and variables > Actions). The workflow reads them as '
     '${{ secrets.DOCKERHUB_USERNAME }} and ${{ secrets.DOCKERHUB_TOKEN }}, and GitHub masks them in logs. A scoped Docker Hub '
     'access token is used instead of the account password so it can be revoked on its own.'),
    ('Q3. What component collects CPU and memory metrics from kubelets so that kubectl top pods works?',
     "The Kubernetes Metrics Server. It scrapes resource usage from each node's kubelet and exposes it through the metrics.k8s.io "
     'API, which kubectl top and the Horizontal Pod Autoscaler use. In Minikube it is enabled with '
     '"minikube addons enable metrics-server".'),
]
for q, a in qa:
    bold(q); d.add_paragraph(a)

d.add_heading('3. Environment Setup', 1)
img('00_setup_tools.png', 'Fig 1. Installed tool versions')
img('01_minikube.png', 'Fig 2. Minikube cluster running, node Ready')

d.add_heading('4. Activity 1: Deployment and Replica Scaling (CO3)', 1)
d.add_paragraph('deployment.yaml:')
code(open(os.path.join(here, 'deployment.yaml')).read())
d.add_paragraph('Commands: kubectl apply -f deployment.yaml; kubectl get deployments; kubectl get pods -l app=web; '
                'kubectl scale deployment sample-web-app --replicas=5; kubectl get pods -w')
img('04_activity1_scaled_5_pods.png', 'Screenshot 1. kubectl get pods showing 5 active scaled pods')

d.add_heading('5. Activity 2: CI/CD Pipeline with GitHub Actions (CO5)', 1)
d.add_paragraph('Repository secrets DOCKERHUB_USERNAME and DOCKERHUB_TOKEN were added under Settings > Secrets and variables > Actions.')
d.add_paragraph('Dockerfile:'); code(open(os.path.join(here, 'Dockerfile')).read())
d.add_paragraph('.github/workflows/ci-cd.yml:'); code(open(os.path.join(here, '.github/workflows/ci-cd.yml')).read())
img('12_git_push.png', 'Fig. git commit and push to main')
img('14_actions_green_run.png', 'Screenshot 2. Completed green run in GitHub Actions workflow history')
img('16_actions_job_steps.png', 'Fig. Pipeline log: every build-and-push job step succeeded (run #1)')
img('17_dockerhub_image.png', 'Fig. Image manishraje28/sample-web-app:v1.0.1 delivered to Docker Hub')
img('18_deploy_ci_image.png', 'Fig. CI-built image rolled out to all 5 replicas on Minikube')

d.add_heading('6. Activity 3: Cluster and Pod Monitoring (CO6)', 1)
d.add_paragraph('Both Option A (Minikube Dashboard) and Option B (Prometheus and Grafana via Helm) were completed.')
img('05_metrics_server_top.png', 'Fig. metrics-server enabled; kubectl top nodes and pods')
img('06_dashboard_workloads.png', 'Fig. Minikube Kubernetes Dashboard: sample-web-app 5/5 pods')
img('09_helm_prometheus_install.png', 'Fig. kube-prometheus-stack installed with Helm')
img('10_grafana_cpu_usage.png', 'Screenshot 3. Live CPU usage per pod in Grafana')
img('11_grafana_memory_usage.png', 'Fig. Live memory usage per pod in Grafana')

d.add_heading('7. Reflection Questions', 1)
bold('R1. Pipeline Integration: How can ci-cd.yml run kubectl apply against the cluster after pushing a new image tag?')
d.add_paragraph('Add a deploy job that runs after build-and-push (needs: build-and-push). Store the cluster kubeconfig as a secret '
                '(KUBE_CONFIG), write it to ~/.kube/config, then apply the manifest and set the new image. A local Minikube cluster '
                'cannot be reached from GitHub-hosted runners, so this needs either a cloud cluster (EKS, GKE, AKS) or a self-hosted '
                'runner on the machine running Minikube (runs-on: self-hosted). Example:')
code('''  deploy:
    needs: build-and-push
    runs-on: self-hosted        # or ubuntu-latest for a cloud cluster
    steps:
      - uses: actions/checkout@v3
      - uses: azure/setup-kubectl@v3
      - name: Configure kubeconfig
        run: |
          mkdir -p ~/.kube
          echo "${{ secrets.KUBE_CONFIG }}" | base64 -d > ~/.kube/config
      - name: Deploy new image tag
        run: |
          kubectl apply -f deployment.yaml
          kubectl set image deployment/sample-web-app \\
            web-container=${{ secrets.DOCKERHUB_USERNAME }}/sample-web-app:v1.0.${{ github.run_number }}
          kubectl rollout status deployment/sample-web-app --timeout=120s''')
d.add_paragraph('The rollout status step fails the pipeline if the new pods never become ready, and a "kubectl rollout undo" step '
                'can be added for recovery. A GitOps tool such as Argo CD is another option: the pipeline commits the new tag to '
                'the manifest and Argo CD syncs the cluster.')
bold('R2. Resource Constraints: What happens when a container exceeds limits.memory? How would Grafana signal this?')
d.add_paragraph('Memory cannot be throttled the way CPU can. When a container goes over its memory limit (256Mi here), the kernel '
                'OOM killer ends the process. Kubernetes records the container as Terminated with reason OOMKilled (exit code 137) '
                'and restarts it according to restartPolicy. Repeated kills raise the restart count and the pod enters '
                'CrashLoopBackOff. Going over the CPU limit (250m) does not kill the container; it is throttled instead.')
d.add_paragraph('In Grafana with kube-prometheus-stack, the Compute Resources dashboards show memory working set climbing to the '
                'limit line, then dropping as the container restarts. The alert rules bundled with the stack (for example '
                'KubePodCrashLooping) fire through Alertmanager, and a custom rule can catch it directly:')
code('''kube_pod_container_status_last_terminated_reason{reason="OOMKilled"} == 1
or
container_memory_working_set_bytes / on(pod, container)
  kube_pod_container_resource_limits{resource="memory"} > 0.9''')
d.add_paragraph('A Grafana alert on the memory panel can also notify by email or Slack when usage stays above 90% of the limit.')

d.add_heading('8. Conclusion', 1)
d.add_paragraph('A declarative Deployment was created on Minikube and scaled from 2 to 5 replicas. A GitHub Actions pipeline built '
                'the Docker image on a push to main and pushed it to Docker Hub using repository secrets; the first run passed and '
                'the resulting v1.0.1 image was rolled out to the cluster. Cluster and pod health were monitored with '
                'metrics-server, the Kubernetes Dashboard, and Prometheus with Grafana.')
d.save(os.path.join(here, 'Experiment8_Report.docx'))
print('ok')
