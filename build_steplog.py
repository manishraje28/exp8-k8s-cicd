import json, os
from docx import Document
from docx.shared import Inches, Pt
here = os.path.dirname(os.path.abspath(__file__))
d = Document()
d.add_heading('Experiment 8 - Step-by-Step Execution Log', 0)
d.add_paragraph('Deploying Applications, CI/CD Pipeline and Cluster Monitoring (Minikube on Windows 11)')
for i, (txt, img) in enumerate(json.load(open(os.path.join(here, 'steps.json'), encoding='utf-8')), 1):
    p = d.add_paragraph(); r = p.add_run(f'Step {i}: '); r.bold = True; p.add_run(txt)
    path = os.path.join(here, img) if img else ''
    if img and os.path.exists(path): d.add_picture(path, width=Inches(6.5))
d.save(os.path.join(here, 'Experiment8_Step_Log.docx'))
print('saved')
